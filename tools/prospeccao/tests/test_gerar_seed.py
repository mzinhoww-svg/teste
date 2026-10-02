from openpyxl import load_workbook

from central.seed import doc_lead, doc_teste, perfil_lead, seed
from msg.gerar import exportar_planilha, gerar
from msg.prep import carregar


def test_gerar_monta_tres_toques_por_canal(arquivos):
    lp, pp = arquivos
    linhas, erros = gerar(lp, pp)
    assert erros == []
    assert [x["id"] for x in linhas] == ["R0001", "R0002"]  # sem canal fica de fora
    whats, email = linhas
    assert whats["canal"] == "WhatsApp" and email["canal"] == "E-mail"
    assert [t["n"] for t in whats["toques"]] == [1, 2, 3]
    assert all(t["waLink"] and not t["corpo"] for t in whats["toques"])
    assert all(t["assunto"] and t["corpo"] and not t["waLink"] for t in email["toques"])
    # Foto só no WhatsApp do toque 1
    assert whats["foto"] == "mesa-pessoa-02" and "Te mandei uma foto" in whats["toques"][0]["mensagem"]
    assert "foto" not in whats["toques"][1]["mensagem"]
    assert email["foto"] == "" and "foto" not in email["toques"][0]["corpo"]


def test_gerar_acusa_lead_sem_personal(arquivos, tmp_path):
    lp, _ = arquivos
    vazio = tmp_path / "vazio.json"
    vazio.write_text("[]", encoding="utf-8")
    _, erros = gerar(lp, str(vazio))
    assert "R0001: sem entrada em personal.json" in erros


def test_seed_estado_inicial_e_lotes(arquivos):
    linhas, _ = gerar(*arquivos)
    d = doc_lead(linhas[0])
    assert d["etapa"] == 0 and d["situacao"] == "ativo" and d["enviado1"] is None
    assert len(d["toques"]) == 3
    lotes = seed(linhas * 30)
    assert [len(x) for x in lotes] == [50, 10]
    assert doc_teste()["telefone"] == "5565999207108"


def test_planilha_tem_leads_e_copy(arquivos, tmp_path):
    lp, pp = arquivos
    linhas, _ = gerar(lp, pp)
    destino = tmp_path / "saida.xlsx"
    exportar_planilha(str(destino), carregar(lp), linhas)
    wb = load_workbook(destino)
    assert wb.sheetnames == ["Leads", "Copy", "Enriquecimento", "Copy pós-venda"]
    ws = wb["Leads"]
    assert ws.max_row == 4
    col = [c.value for c in ws[1]].index("Status") + 1
    status = [ws.cell(r, col).value for r in range(2, 5)]
    assert status == ["Pendente", "Pendente", "Sem canal"]


def test_base_semeia_config_do_pos_venda():
    from central.seed import base
    ids = [w["doc_id"] for w in base()]
    assert ids == ["TESTE", "meta", "posvenda", "fotos"]


def test_email_que_ganhou_whatsapp_do_decisor_vira_whatsapp(arquivos, tmp_path):
    import json
    lp, pp = arquivos
    enr = [{"id": "R0002", "empresa": {"cnpj": ""}, "socios": [], "decisores": [], "redes": {}, "sinais": [],
            "pendencias": [], "observacao": "",
            "contatos": [{"id": "k1", "papel": "decisor", "nome": "MARIANA SOUZA", "cargo": "Diretora",
                          "telefone": "5565988887777", "whatsapp": "sim", "email": "", "fonte": "site", "confianca": "alta"}]}]
    ep = tmp_path / "enr.json"
    ep.write_text(json.dumps(enr), encoding="utf-8")
    linhas, erros = gerar(lp, pp, str(ep))
    assert erros == []
    x = [l for l in linhas if l["id"] == "R0002"][0]
    assert x["canal"] == "WhatsApp" and x["telefone"] == "5565988887777"
    assert x["toques"][0]["mensagem"].startswith("Oi, Mariana, tudo bem?")
    assert "WhatsApp do enriquecimento" in x["flags"]


def test_perfil_do_lead_sem_telefone_nem_observacao():
    lead = {"id": "R0001", "especialidade": "dermatologia clínica", "porte": "pequeno", "cidade": "Cuiabá",
            "nota": 4.9, "avaliacoes": 120, "fonte": "https://exemplo.com.br/", "obs": "WhatsApp (65) 99999-1111",
            "telefone": "5565999991111"}
    p = perfil_lead(lead, {"fonte": "Especialidade"})
    assert p == {"especialidade": "dermatologia clínica", "porte": "pequeno", "cidade": "Cuiabá", "nota": 4.9,
                 "avaliacoes": 120, "fonteDados": "https://exemplo.com.br/", "fonteFrase": "Especialidade"}
    assert perfil_lead({"id": "R2"}, {})["especialidade"] == ""
