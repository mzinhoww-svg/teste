"""Gera as mensagens dos três toques, checa tudo e exporta a planilha.

Uso:
  python3 -m msg.gerar                      # checagem; precisa dar "0 erros"
  python3 -m msg.gerar --planilha Reiners_Leads_Cuiaba.xlsx
"""
import argparse
import json
from collections import Counter
from dataclasses import replace

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from msg.checks import Personal, check_lote, check_personal, check_toque
from msg.compose import compose_email, compose_whatsapp, wa_link
from msg.copy_posvenda import linhas_copy as linhas_copy_pv
from msg.enriquecimento import PAPEIS, ROTULO_PAPEL, carregar as carregar_enriq, contato_direto, primeiro_nome
from msg.enriquecimento import status as status_enriq
from msg.fotos import CATALOGO, foto_para, linha_foto
from msg.copy_v1 import ICPS, TOQUES, VERSAO, linhas_copy
from msg.prep import Lead, canal, carregar


def carregar_personal(caminho: str) -> dict[str, Personal]:
    with open(caminho, encoding="utf-8") as fh:
        return {d["id"]: Personal(**d) for d in json.load(fh)}


def montar(l: Lead, p: Personal, ordem: int) -> dict:
    c = canal(l)
    foto = foto_para(l.icp, l.nome, l.categoria, l.especialidade)
    toques = []
    for t in TOQUES:
        # A foto vai junto só no WhatsApp do toque 1; o e-mail não leva anexo.
        lf = linha_foto(foto) if (t == 1 and c == "WhatsApp") else ""
        texto = compose_whatsapp(t, l.icp, p.saudacao, p.nome_curto, p.frase, lf)
        item = {"n": t, "mensagem": "", "waLink": "", "assunto": "", "corpo": "", "_texto": texto, "_linha_foto": lf}
        if c == "WhatsApp":
            item["mensagem"] = texto
            item["waLink"] = wa_link(l.telefone, texto)
        else:
            item["assunto"], item["corpo"] = compose_email(t, l.icp, p.saudacao, p.nome_curto, p.frase)
        toques.append(item)
    return {
        "id": l.id, "ordem": ordem, "nome": l.nome, "icp": l.icp, "segmento": ICPS[l.icp],
        "categoria": l.categoria, "faixa": l.faixa, "score": l.score, "bairro": l.bairro,
        "canal": c, "telefone": l.telefone, "email": l.email, "site": l.site,
        "instagram": l.instagram, "saudacao": p.saudacao, "nomeCurto": p.nome_curto,
        "fraseUnica": p.frase, "fonte": p.fonte, "flags": list(l.flags), "versaoCopy": VERSAO,
        "foto": foto if c == "WhatsApp" else "",
        "toques": toques,
    }


def aplicar_enriquecimento(l: Lead, p: Personal, enr: dict | None) -> tuple[Lead, Personal]:
    """Lead só com e-mail que ganhou WhatsApp de quem decide (ou da comunicação) vira lead de WhatsApp,
    com a saudação pelo primeiro nome da pessoa. Os demais seguem como estão: a troca de contato de
    um lead de WhatsApp é feita no card da central ("Usar na cadência")."""
    if not enr or canal(l) != "E-mail":
        return l, p
    d = contato_direto(enr)
    if not d or not d["telefone"] or d["whatsapp"] != "sim":
        return l, p
    l = replace(l, telefone=d["telefone"], celular=d["telefone"][4] == "9", whatsapp_fixo=d["telefone"][4] != "9",
                flags=l.flags + ["WhatsApp do enriquecimento"])
    nome = primeiro_nome(d["nome"])
    if nome:
        p = replace(p, saudacao=nome)
    return l, p


def gerar(leads_path: str, personal_path: str, enriq_path: str | None = None) -> tuple[list[dict], list[str]]:
    """Devolve (linhas, erros) só para leads com canal."""
    leads = [l for l in carregar(leads_path) if canal(l) != "Sem canal"]
    personal = carregar_personal(personal_path)
    enriq = carregar_enriq(enriq_path) if enriq_path else {}
    linhas, erros, pares = [], [], []
    for ordem, l in enumerate(leads, start=1):
        p = personal.get(l.id)
        if p is None:
            erros.append(f"{l.id}: sem entrada em personal.json")
            continue
        l, p = aplicar_enriquecimento(l, p, enriq.get(l.id))
        erros += check_personal(l, p)
        linha = montar(l, p, ordem)
        for t in linha["toques"]:
            erros += check_toque(l, p, t["n"], t.pop("_texto"), t["waLink"], t.pop("_linha_foto"))
        linhas.append(linha)
        pares.append((l, p))
    erros += check_lote(pares)
    return linhas, erros


COLUNAS = [
    ("ID", "id", 8), ("ICP", "icp", 7), ("Segmento", "segmento", 18), ("Nome", "nome", 32),
    ("Categoria", "categoria", 22), ("Bairro", "bairro", 16), ("Cidade", "cidade", 12),
    ("Telefone", "telefone", 15), ("Celular", "celular", 8), ("Canal", "canal", 10),
    ("E-mail", "email", 26), ("Site", "site", 28), ("Instagram", "instagram", 22),
    ("Responsável", "responsavel", 18), ("Especialidade", "especialidade", 40),
    ("Porte", "porte", 9), ("Score", "score", 7), ("Faixa", "faixa", 6), ("Flags", "flags", 16),
    ("Fonte do dado", "fonte", 30), ("Observação", "obs", 30),
    ("Saudação", "saudacao", 22), ("Nome curto", "nomeCurto", 18), ("Frase única", "fraseUnica", 50),
    ("Foto do toque 1", "foto", 22), ("Toque 1", "t1", 60), ("Toque 2", "t2", 60), ("Toque 3", "t3", 60),
    ("Status", "status", 12),
    ("CNPJ", "cnpj", 16), ("Razão social", "razao", 30), ("Quem lidera", "lider", 24), ("Cargo", "cargoLider", 18),
    ("Contato sugerido", "sugerido", 34), ("Enriquecimento", "statusEnriq", 12),
]


def exportar_planilha(destino: str, leads: list[Lead], linhas: list[dict], enriq: dict | None = None) -> None:
    por_id = {x["id"]: x for x in linhas}
    enriq = enriq or {}
    wb = Workbook()
    ws = wb.active
    ws.title = "Leads"
    cab = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    fundo = PatternFill("solid", fgColor="0F0F0F")
    corpo = Font(name="Arial", size=10)
    ws.append([c[0] for c in COLUNAS])
    for c in ws[1]:
        c.font, c.fill = cab, fundo
    for l in leads:
        d = l.to_dict()
        d["canal"] = canal(l)
        d["segmento"] = ICPS.get(l.icp, "")
        d["celular"] = "Sim" if l.celular else ("WhatsApp fixo" if l.whatsapp_fixo else "Não")
        d["flags"] = ", ".join(l.flags)
        x = por_id.get(l.id)
        if x:
            d.update(saudacao=x["saudacao"], nomeCurto=x["nomeCurto"], fraseUnica=x["fraseUnica"],
                     foto=f"{x['foto']}.jpg ({CATALOGO[x['foto']][0]})" if x["foto"] else "")
            for t in x["toques"]:
                d[f"t{t['n']}"] = t["mensagem"] or f"Assunto: {t['assunto']}\n\n{t['corpo']}"
            d["status"] = "Pendente"
        else:
            d["status"] = "Sem canal"
        e = enriq.get(l.id)
        if e:
            d["cnpj"], d["razao"] = e["empresa"].get("cnpj", ""), e["empresa"].get("razaoSocial", "")
            if e["decisores"]:
                d["lider"], d["cargoLider"] = e["decisores"][0]["nome"], e["decisores"][0]["cargo"]
            s_ = contato_direto(e)
            if s_:
                d["sugerido"] = " · ".join(x for x in (s_["nome"], ROTULO_PAPEL[s_["papel"]],
                                                       ("+" + s_["telefone"]) if s_["telefone"] else s_["email"]) if x)
            d["statusEnriq"] = status_enriq(e)
        else:
            d["statusEnriq"] = "bruto"
        ws.append([d.get(k, "") if d.get(k) is not None else "" for _, k, _ in COLUNAS])
    for i, (_, _, w) in enumerate(COLUNAS, start=1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = corpo
            c.alignment = Alignment(wrap_text=c.column_letter in ("O", "X", "Z", "AA", "AB"), vertical="top")
    ws.freeze_panes = "E2"

    cp = wb.create_sheet("Copy")
    cp.append(["Bloco", "ICP", "Texto"])
    for c in cp[1]:
        c.font, c.fill = cab, fundo
    for linha in linhas_copy():
        cp.append(list(linha))
    for row in cp.iter_rows(min_row=2):
        for c in row:
            c.font = corpo
            c.alignment = Alignment(wrap_text=True, vertical="top")
    cp.column_dimensions["A"].width = 30
    cp.column_dimensions["B"].width = 8
    cp.column_dimensions["C"].width = 100

    en = wb.create_sheet("Enriquecimento")
    en.append(["ID", "Lead", "Papel", "Nome", "Cargo", "Telefone", "WhatsApp", "E-mail", "Fonte", "Confiança"])
    for c in en[1]:
        c.font, c.fill = cab, fundo
    nomes = {l.id: l.nome for l in leads}
    for lid, e in enriq.items():
        for d_ in e["decisores"]:
            en.append([lid, nomes.get(lid, ""), "Quem lidera", d_["nome"], d_["cargo"], "", "", "", d_["fonte"], d_["confianca"]])
        for c_ in e["contatos"]:
            en.append([lid, nomes.get(lid, ""), ROTULO_PAPEL[c_["papel"]], c_["nome"], c_["cargo"],
                       ("+" + c_["telefone"]) if c_["telefone"] else "", c_["whatsapp"], c_["email"], c_["fonte"], c_["confianca"]])
    for row in en.iter_rows(min_row=2):
        for c in row:
            c.font = corpo
    for col, w in zip("ABCDEFGHIJ", (8, 30, 14, 26, 22, 17, 9, 28, 40, 10)):
        en.column_dimensions[col].width = w

    pv = wb.create_sheet("Copy pós-venda")
    pv.append(["Etapa", "Quando / variante", "Texto"])
    for c in pv[1]:
        c.font, c.fill = cab, fundo
    for linha in linhas_copy_pv():
        pv.append(list(linha))
    for row in pv.iter_rows(min_row=2):
        for c in row:
            c.font = corpo
            c.alignment = Alignment(wrap_text=True, vertical="top")
    pv.column_dimensions["A"].width = 24
    pv.column_dimensions["B"].width = 28
    pv.column_dimensions["C"].width = 100
    wb.save(destino)


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--leads", default="dados/leads.json")
    ap.add_argument("--personal", default="msg/personal.json")
    ap.add_argument("--planilha", help="exporta a planilha .xlsx com leads, mensagens e copy")
    ap.add_argument("--out", help="grava as mensagens geradas neste JSON")
    ap.add_argument("--enriq", default="dados/enriquecimento.json", help="registros de enriquecimento validados")
    a = ap.parse_args()
    linhas, erros = gerar(a.leads, a.personal, a.enriq)
    print("\n".join(erros) or "0 erros")
    print(f"{len(linhas)} leads com mensagem")
    print("Canais:", dict(Counter(x["canal"] for x in linhas)))
    print("Fontes:", dict(Counter(x["fonte"] for x in linhas)))
    maior = max((len(t["mensagem"]) for x in linhas for t in x["toques"]), default=0)
    print("Maior WhatsApp:", maior, "caracteres")
    if a.out:
        with open(a.out, "w", encoding="utf-8") as fh:
            json.dump(linhas, fh, ensure_ascii=False, indent=1)
    if a.planilha:
        exportar_planilha(a.planilha, carregar(a.leads), linhas, carregar_enriq(a.enriq))
        print("Planilha:", a.planilha)


if __name__ == "__main__":
    main()
